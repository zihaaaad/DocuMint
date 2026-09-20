import os
import sys
import pandas as pd
from docx import Document

# Ensure src is in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'src')))

from documint.core import (
    is_valid_email,
    apply_template_replacements,
    _replace_in_paragraph,
    _replace_in_table,
    replace_placeholders_in_doc,
    validate_placeholders,
    load_dataset,
    SMTPSender
)

def test_apply_template_replacements():
    """Verify unified string replacement handles <var>, {{var}}, [var], and {var}."""
    replacements = {"Name": "Zihad Hasan", "ID": "12345", "Course": "AI Systems"}
    
    s1 = "Hello <Name>, ID: <ID>."
    assert apply_template_replacements(s1, replacements) == "Hello Zihad Hasan, ID: 12345."

    s2 = "Student {{Name}} enrolled in {{Course}}."
    assert apply_template_replacements(s2, replacements) == "Student Zihad Hasan enrolled in AI Systems."

    s3 = "Result for [Name]: Score is [ID]"
    assert apply_template_replacements(s3, replacements) == "Result for Zihad Hasan: Score is 12345"

    s4 = "Filename_{Name}_{ID}"
    assert apply_template_replacements(s4, replacements) == "Filename_Zihad Hasan_12345"

def test_email_validation():
    """Test valid and invalid email addresses."""
    assert is_valid_email("user@example.com") is True
    assert is_valid_email("first.last@domain.co.uk") is True
    assert is_valid_email("student+tag@university.edu") is True
    assert is_valid_email("invalid-email") is False
    assert is_valid_email("") is False
    assert is_valid_email("@domain.com") is False
    assert is_valid_email(None) is False

def test_replace_placeholders_in_body_and_tables():
    """Test placeholder replacement across paragraphs and multi-cell tables."""
    doc = Document()
    p1 = doc.add_paragraph("Welcome <Name>, to {{Course}}.")
    p2 = doc.add_paragraph("Your score is [Score].")
    
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].paragraphs[0].text = "Header <Dept>"
    table.rows[0].cells[1].paragraphs[0].text = "Value {{Score}}"
    table.rows[1].cells[0].paragraphs[0].text = "Student: [Name]"
    table.rows[1].cells[1].paragraphs[0].text = "Status: Active"

    replacements = {
        "Name": "Zihad Hasan",
        "Course": "Deep Learning",
        "Score": "98%",
        "Dept": "Computer Science"
    }
    
    replace_placeholders_in_doc(doc, replacements)
    
    assert doc.paragraphs[0].text == "Welcome Zihad Hasan, to Deep Learning."
    assert doc.paragraphs[1].text == "Your score is 98%."
    assert table.rows[0].cells[0].paragraphs[0].text == "Header Computer Science"
    assert table.rows[0].cells[1].paragraphs[0].text == "Value 98%"
    assert table.rows[1].cells[0].paragraphs[0].text == "Student: Zihad Hasan"

def test_replace_placeholders_in_headers_and_footers():
    """Verify that headers and footers in document sections have placeholders replaced."""
    doc = Document()
    section = doc.sections[0]
    
    header = section.header
    header.paragraphs[0].text = "Institution: <Institution> | Year: {{Year}}"
    
    footer = section.footer
    footer.paragraphs[0].text = "Confidential - [Doc_ID]"

    replacements = {
        "Institution": "MIT",
        "Year": "2026",
        "Doc_ID": "DOC-9941"
    }

    replace_placeholders_in_doc(doc, replacements)

    assert header.paragraphs[0].text == "Institution: MIT | Year: 2026"
    assert footer.paragraphs[0].text == "Confidential - DOC-9941"

def test_run_formatting_preservation():
    """Verify that replacing a placeholder inside a formatted run preserves character styling."""
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("<Participant>")
    run.bold = True
    run.italic = True
    
    replacements = {"Participant": "Zihad Hasan"}
    replace_placeholders_in_doc(doc, replacements)
    
    assert p.text == "Zihad Hasan"
    assert p.runs[0].bold is True
    assert p.runs[0].italic is True

def test_load_dataset(tmp_path):
    """Verify loading from both Excel and CSV formats."""
    # CSV
    csv_file = tmp_path / "data.csv"
    df_src = pd.DataFrame({"Name ": ["Alice", "Bob"], " Email": ["a@test.com", "b@test.com"]})
    df_src.to_csv(csv_file, index=False)
    
    df_loaded = load_dataset(str(csv_file))
    assert list(df_loaded.columns) == ["Name", "Email"]
    assert len(df_loaded) == 2

    # Excel
    xlsx_file = tmp_path / "data.xlsx"
    df_src.to_excel(xlsx_file, index=False)
    df_xlsx = load_dataset(str(xlsx_file))
    assert list(df_xlsx.columns) == ["Name", "Email"]
    assert len(df_xlsx) == 2

def test_validate_placeholders(tmp_path):
    """Verify placeholder validation against dataset columns."""
    csv_file = tmp_path / "data.csv"
    df = pd.DataFrame({"Name": ["Zihad"], "Score": [99], "Institution": ["MIT"]})
    df.to_csv(csv_file, index=False)
    
    doc_file = tmp_path / "template.docx"
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Header <Institution>"
    doc.add_paragraph("Hello <Name>, your score is {{Score}}.")
    doc.save(str(doc_file))
    
    is_valid, missing = validate_placeholders(str(csv_file), str(doc_file))
    assert is_valid is True
    assert missing == []
    
    # Missing variable in template
    doc_missing = tmp_path / "template_missing.docx"
    doc2 = Document()
    doc2.add_paragraph("Hello <Name>, your score is {{Score}} for course <MissingCourse>.")
    doc2.save(str(doc_missing))
    
    is_valid_bad, missing_bad = validate_placeholders(str(csv_file), str(doc_missing))
    assert is_valid_bad is False
    assert "MissingCourse" in missing_bad

def test_smtp_sender_port_configuration():
    """Verify SMTPSender configures SSL on port 465 and STARTTLS on port 587."""
    ssl_sender = SMTPSender(host="smtp.gmail.com", port=465, user="test@gmail.com", password="pwd")
    assert ssl_sender.use_ssl is True
    assert ssl_sender.use_starttls is False

    tls_sender = SMTPSender(host="smtp.gmail.com", port=587, user="test@gmail.com", password="pwd")
    assert tls_sender.use_ssl is False
    assert tls_sender.use_starttls is True

def test_process_emails_dry_run(tmp_path):
    """Verify process_emails executes cleanly in dry-run mode without email dispatch."""
    csv_file = tmp_path / "data.csv"
    pd.DataFrame({
        "Name": ["Alice", "Bob"],
        "Email": ["alice@example.com", "bob@example.com"]
    }).to_csv(csv_file, index=False)

    doc_file = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("Hello <Name>")
    doc.save(str(doc_file))

    pdf_out = tmp_path / "pdf"
    logs_out = tmp_path / "logs"

    from documint.core import process_emails

    summary = process_emails(
        data_file=str(csv_file),
        template_file=str(doc_file),
        pdf_folder=str(pdf_out),
        logs_folder=str(logs_out),
        email_subject="Notice for <Name>",
        email_body="<p>Dear <Name></p>",
        pdf_filename_format="Cert_{Name}",
        dry_run=True
    )

    assert summary["total"] == 2
    assert summary["success"] == 2
    assert summary["failed"] == 0

