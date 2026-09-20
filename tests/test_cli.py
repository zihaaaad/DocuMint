import os
import sys
import pytest
from docx import Document
import pandas as pd

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'src')))

from documint.cli import main

def test_cli_validate_success(tmp_path, capsys):
    csv_file = tmp_path / "data.csv"
    pd.DataFrame({"Name": ["Alice"], "ID": ["101"]}).to_csv(csv_file, index=False)

    doc_file = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("Hello <Name>, your ID is <ID>.")
    doc.save(str(doc_file))

    with pytest.raises(SystemExit) as exc_info:
        main(["validate", "--data", str(csv_file), "--template", str(doc_file)])
    
    assert exc_info.value.code == 0
    captured = capsys.readouterr()
    assert "Pre-Flight Validation PASSED" in captured.out

def test_cli_validate_failure(tmp_path, capsys):
    csv_file = tmp_path / "data.csv"
    pd.DataFrame({"Name": ["Alice"]}).to_csv(csv_file, index=False)

    doc_file = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("Hello <Name>, your ID is <MissingID>.")
    doc.save(str(doc_file))

    with pytest.raises(SystemExit) as exc_info:
        main(["validate", "--data", str(csv_file), "--template", str(doc_file)])
    
    assert exc_info.value.code == 1
    captured = capsys.readouterr()
    assert "Pre-Flight Validation FAILED" in captured.out
