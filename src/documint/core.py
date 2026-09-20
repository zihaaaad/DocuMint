import os
import re
import time
import shutil
import logging
import smtplib
import mimetypes
import threading
import concurrent.futures
from datetime import datetime
from abc import ABC, abstractmethod
from typing import Dict, Any, Optional, Callable, List, Tuple

import pandas as pd
from docx import Document
from email.message import EmailMessage

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# --- Abstract Base Classes (The "Universal" Layer) ---

class DocumentConverter(ABC):
    """Abstract base class for converting documents to PDF."""
    
    @abstractmethod
    def convert_to_pdf(self, input_path: str, output_path: str) -> None:
        """Converts a source document to PDF."""
        pass

class EmailSender(ABC):
    """Abstract base class for sending emails."""
    
    @abstractmethod
    def send_email(self, to_email: str, subject: str, body_html: str, attachments: Optional[List[str]] = None) -> None:
        """Sends an email with optional attachments."""
        pass

# --- Universal Document Converter Cascade ---

class UniversalDocumentConverter(DocumentConverter):
    """Universal converter that dynamically selects the best available engine (Word COM, docx2pdf, or LibreOffice)."""
    
    def __init__(self):
        self._word_app = None
        self._is_com_initialized = False

    def __enter__(self):
        self._init_session()
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()

    def _init_session(self):
        if os.name == 'nt' and self._word_app is None:
            try:
                import win32com.client
                import pythoncom
                pythoncom.CoInitialize()
                self._is_com_initialized = True
                self._word_app = win32com.client.DispatchEx("Word.Application")
                self._word_app.Visible = False
                self._word_app.DisplayAlerts = 0
            except Exception as e:
                logger.warning(f"Could not initialize persistent Word COM session: {e}")
                self._word_app = None

    def close(self):
        if self._word_app is not None:
            try:
                self._word_app.Quit(0)  # 0 = wdDoNotSaveChanges
            except Exception:
                pass
            finally:
                self._word_app = None
        if self._is_com_initialized:
            try:
                import pythoncom
                pythoncom.CoUninitialize()
            except Exception:
                pass
            self._is_com_initialized = False

    def convert_to_pdf(self, input_path: str, output_path: str) -> None:
        abs_input = os.path.abspath(input_path)
        abs_output = os.path.abspath(output_path)
        output_dir = os.path.dirname(abs_output)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir, exist_ok=True)
        
        # 1. Try Windows Word COM (Reusing persistent session if active)
        if os.name == 'nt':
            created_local = False
            word = self._word_app
            try:
                if word is None:
                    import win32com.client
                    import pythoncom
                    pythoncom.CoInitialize()
                    word = win32com.client.DispatchEx("Word.Application")
                    word.Visible = False
                    word.DisplayAlerts = 0
                    created_local = True

                doc = None
                try:
                    doc = word.Documents.Open(abs_input)
                    doc.SaveAs(abs_output, FileFormat=17)  # 17 = wdFormatPDF
                finally:
                    if doc is not None:
                        try:
                            doc.Close(0)
                        except Exception:
                            pass
                        doc = None
                if os.path.exists(abs_output):
                    return
            except Exception as e:
                logger.warning(f"WinWord COM conversion error: {e}. Attempting fallbacks...")
            finally:
                if created_local and word is not None:
                    try:
                        word.Quit(0)
                    except Exception:
                        pass
                    word = None
                    try:
                        import pythoncom
                        pythoncom.CoUninitialize()
                    except Exception:
                        pass

        # 2. Try docx2pdf
        try:
            from docx2pdf import convert
            convert(abs_input, abs_output)
            if os.path.exists(abs_output):
                return
        except Exception:
            pass

        # 3. Try LibreOffice CLI
        import subprocess
        possible_cmds = ['libreoffice', 'soffice']
        if os.name == 'nt':
            # Check standard Windows LibreOffice installation directories
            for lo_path in [
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                os.path.expandvars(r"%LOCALAPPDATA%\Programs\LibreOffice\program\soffice.exe")
            ]:
                if os.path.exists(lo_path):
                    possible_cmds.insert(0, lo_path)

        for cmd in possible_cmds:
            try:
                subprocess.run(
                    [cmd, '--headless', '--convert-to', 'pdf', abs_input, '--outdir', output_dir],
                    check=True,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    timeout=45
                )
                generated_pdf = os.path.join(output_dir, os.path.splitext(os.path.basename(abs_input))[0] + '.pdf')
                if os.path.exists(generated_pdf) and os.path.abspath(generated_pdf) != abs_output:
                    shutil.move(generated_pdf, abs_output)
                if os.path.exists(abs_output):
                    return
            except Exception:
                continue

        raise RuntimeError(
            "No PDF conversion engine available. Please ensure Microsoft Word (Windows) or LibreOffice is installed."
        )

class WinWordConverter(UniversalDocumentConverter):
    """Compatibility alias for UniversalDocumentConverter."""
    pass

# --- Email Dispatch Implementations ---

class WinOutlookSender(EmailSender):
    """Uses Microsoft Outlook (via COM) to send emails."""
    
    def __init__(self):
        import win32com.client
        self.outlook = win32com.client.Dispatch("Outlook.Application")

    def send_email(self, to_email: str, subject: str, body_html: str, attachments: Optional[List[str]] = None) -> None:
        try:
            mail = self.outlook.CreateItem(0)
            mail.To = to_email
            mail.Subject = subject
            mail.HTMLBody = body_html
            
            if attachments:
                for att in attachments:
                    if os.path.exists(att):
                        mail.Attachments.Add(os.path.abspath(att))
            
            mail.Send()
        except Exception as e:
            logger.error(f"Outlook send failed: {e}")
            raise e

class SMTPSender(EmailSender):
    """Universal SMTP sender with support for Port 465 (SSL) and Port 587/25 (STARTTLS)."""
    
    def __init__(
        self,
        host: str,
        port: int = 587,
        user: str = "",
        password: str = "",
        use_tls: Optional[bool] = None,
        sender_name: Optional[str] = None
    ):
        self.host = host.strip()
        self.port = int(port)
        self.user = user.strip()
        self.password = password
        self.sender_name = sender_name.strip() if sender_name else None
        
        # Auto-detect TLS if not explicitly specified: Port 465 is SSL, 587/25 is STARTTLS
        if use_tls is None:
            self.use_ssl = (self.port == 465)
            self.use_starttls = (self.port in (587, 25, 2525))
        else:
            self.use_ssl = not use_tls and (self.port == 465)
            self.use_starttls = use_tls or (self.port != 465)

    def send_email(self, to_email: str, subject: str, body_html: str, attachments: Optional[List[str]] = None) -> None:
        msg = EmailMessage()
        msg['Subject'] = subject
        if self.sender_name:
            msg['From'] = f"{self.sender_name} <{self.user}>"
        else:
            msg['From'] = self.user
        msg['To'] = to_email
        msg.set_content("Please enable HTML in your email client to view this message.")
        msg.add_alternative(body_html, subtype='html')

        if attachments:
            for att in attachments:
                if os.path.exists(att):
                    ctype, encoding = mimetypes.guess_type(att)
                    if ctype is None or encoding is not None:
                        ctype = 'application/octet-stream'
                    
                    maintype, subtype = ctype.split('/', 1)
                    with open(att, 'rb') as f:
                        file_data = f.read()
                        msg.add_attachment(
                            file_data,
                            maintype=maintype,
                            subtype=subtype,
                            filename=os.path.basename(att)
                        )

        try:
            if self.use_ssl:
                with smtplib.SMTP_SSL(self.host, self.port, timeout=30) as smtp:
                    if self.user and self.password:
                        smtp.login(self.user, self.password)
                    smtp.send_message(msg)
            else:
                with smtplib.SMTP(self.host, self.port, timeout=30) as smtp:
                    if self.use_starttls:
                        smtp.starttls()
                    if self.user and self.password:
                        smtp.login(self.user, self.password)
                    smtp.send_message(msg)
        except Exception as e:
            logger.error(f"SMTP send failed to {to_email}: {e}")
            raise e

# --- Core Template Replacement & Formatting Engine ---

def _build_expanded_replacements(replacements: Dict[str, Any]) -> Dict[str, str]:
    """Expands a dictionary of key-value pairs into delimited token syntax variants."""
    expanded = {}
    for k, v in replacements.items():
        val_str = str(v)
        raw_k = str(k).strip()
        clean_k = raw_k.strip('<>{}[]')
        
        if raw_k != clean_k:
            expanded[raw_k] = val_str
        
        expanded[f"<{clean_k}>"] = val_str
        expanded[f"{{{{{clean_k}}}}}"] = val_str
        expanded[f"[{clean_k}]"] = val_str
        expanded[f"{{{clean_k}}}"] = val_str
    return expanded

def apply_template_replacements(text: str, replacements: Dict[str, Any]) -> str:
    """Replaces placeholders in a plain text or HTML string across all token syntaxes."""
    if not text:
        return text
    expanded = _build_expanded_replacements(replacements)
    # Sort keys by length in descending order to prevent substring collisions
    sorted_keys = sorted(expanded.keys(), key=len, reverse=True)
    res = text
    for key in sorted_keys:
        if key in res:
            res = res.replace(key, expanded[key])
    return res

def _replace_in_paragraph(paragraph, replacements: Dict[str, Any]) -> None:
    """Replaces placeholders in a docx paragraph while preserving character-level run formatting."""
    expanded = _build_expanded_replacements(replacements)
    full_text = "".join(run.text for run in paragraph.runs)
    if not any(key in full_text for key in expanded):
        return

    sorted_keys = sorted(expanded.keys(), key=len, reverse=True)

    # 1. First attempt: run-by-run replacement to preserve bold/italic/font formatting
    for run in paragraph.runs:
        for key in sorted_keys:
            if key in run.text:
                run.text = run.text.replace(key, expanded[key])

    # 2. Fallback: if placeholder was split across runs by Word editor cursor edits
    remaining_text = "".join(run.text for run in paragraph.runs)
    if any(key in remaining_text for key in sorted_keys):
        new_text = remaining_text
        for key in sorted_keys:
            if key in new_text:
                new_text = new_text.replace(key, expanded[key])
        if new_text != remaining_text:
            p = paragraph._p
            for child in list(p):
                p.remove(child)
            paragraph.add_run(new_text)

def _replace_in_table(table, replacements: Dict[str, Any]) -> None:
    """Replaces placeholders in all cells and nested structures of a docx table."""
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                _replace_in_paragraph(para, replacements)
            for nested_table in cell.tables:
                _replace_in_table(nested_table, replacements)

def replace_placeholders_in_doc(doc: Document, replacements: Dict[str, Any]) -> None:
    """Replaces placeholders throughout an entire Word document including paragraphs, tables, headers, and footers."""
    # 1. Main body paragraphs
    for para in doc.paragraphs:
        _replace_in_paragraph(para, replacements)
    
    # 2. Main body tables
    for table in doc.tables:
        _replace_in_table(table, replacements)

    # 3. Document sections (Headers and Footers)
    for section in doc.sections:
        # Header
        if section.header:
            for para in section.header.paragraphs:
                _replace_in_paragraph(para, replacements)
            for table in section.header.tables:
                _replace_in_table(table, replacements)
        # Footer
        if section.footer:
            for para in section.footer.paragraphs:
                _replace_in_paragraph(para, replacements)
            for table in section.footer.tables:
                _replace_in_table(table, replacements)
        # First page header / footer (if distinct)
        if getattr(section, 'different_first_page_header_footer', False):
            if section.first_page_header:
                for para in section.first_page_header.paragraphs:
                    _replace_in_paragraph(para, replacements)
                for table in section.first_page_header.tables:
                    _replace_in_table(table, replacements)
            if section.first_page_footer:
                for para in section.first_page_footer.paragraphs:
                    _replace_in_paragraph(para, replacements)
                for table in section.first_page_footer.tables:
                    _replace_in_table(table, replacements)
        # Even page header / footer (if distinct)
        if getattr(doc.settings, 'odd_and_even_pages_header_footer', False):
            if section.even_page_header:
                for para in section.even_page_header.paragraphs:
                    _replace_in_paragraph(para, replacements)
                for table in section.even_page_header.tables:
                    _replace_in_table(table, replacements)
            if section.even_page_footer:
                for para in section.even_page_footer.paragraphs:
                    _replace_in_paragraph(para, replacements)
                for table in section.even_page_footer.tables:
                    _replace_in_table(table, replacements)

def is_valid_email(email_address: str) -> bool:
    """Validates an email address syntax."""
    if not email_address or not isinstance(email_address, str):
        return False
    pattern = r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$"
    return re.match(pattern, email_address.strip()) is not None

def load_dataset(data_file: str) -> pd.DataFrame:
    """Loads dataset from Excel (.xlsx, .xls) or CSV (.csv) file."""
    if not os.path.exists(data_file):
        raise FileNotFoundError(f"Dataset file not found: {data_file}")
    
    if data_file.lower().endswith(('.xlsx', '.xls')):
        df = pd.read_excel(data_file)
    elif data_file.lower().endswith(('.csv', '.txt')):
        df = pd.read_csv(data_file)
    else:
        # Fallback attempt excel then csv
        try:
            df = pd.read_excel(data_file)
        except Exception:
            df = pd.read_csv(data_file)
            
    df.columns = df.columns.astype(str).str.strip()
    return df

def validate_placeholders(data_file: str, template_file: str) -> Tuple[bool, List[str]]:
    """
    Checks if all placeholders in the template exist as columns in the data file.
    Supports <Column>, {{Column}}, [Column], and {Column} syntaxes.
    Returns (True, []) if valid, or (False, list_of_missing_columns).
    """
    if not os.path.exists(data_file):
        return False, [f"Data file not found: {data_file}"]
    if not os.path.exists(template_file):
        return False, [f"Template file not found: {template_file}"]

    try:
        df = load_dataset(data_file)
        data_columns = set(df.columns)

        doc = Document(template_file)
        placeholders = set()
        pattern = r"(?:<|{{|\[|{)([^>}\]]+)(?:>|}}|\]|})"
        
        def extract_from_text(text: str):
            if not text:
                return
            for m in re.findall(pattern, text):
                cleaned = m.strip()
                if cleaned:
                    placeholders.add(cleaned)

        # Paragraphs & tables in body
        for para in doc.paragraphs:
            extract_from_text(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        extract_from_text(para.text)

        # Sections (Headers and Footers)
        for section in doc.sections:
            if section.header:
                for para in section.header.paragraphs:
                    extract_from_text(para.text)
            if section.footer:
                for para in section.footer.paragraphs:
                    extract_from_text(para.text)

        missing = [p for p in placeholders if p not in data_columns]
        return (len(missing) == 0), missing

    except Exception as e:
        logger.error(f"Validation failed: {e}")
        return False, [str(e)]

# --- High-Throughput Batch Processing Engine ---

def process_emails(
    data_file: str, 
    template_file: str, 
    pdf_folder: str, 
    logs_folder: str, 
    log_callback: Callable[[str], None] = lambda msg: None, 
    email_subject: str = "", 
    email_body: str = "", 
    pdf_filename_format: str = "{Name}", 
    retries: int = 1, 
    delay: int = 0, 
    dry_run: bool = False, 
    test_email: Optional[str] = None,
    email_config: Optional[Dict[str, Any]] = None,
    cancel_event: Optional[threading.Event] = None,
    db_manager: Optional[Any] = None
) -> Dict[str, Any]:
    """
    Executes universal batch document compilation and multi-channel email dispatch.
    
    Returns a summary dictionary with counts and itemized logs.
    """
    os.makedirs(pdf_folder, exist_ok=True)
    os.makedirs(logs_folder, exist_ok=True)
    log_file_path = os.path.join(logs_folder, "documint_log.xlsx")
    log_records = []

    # 1. Initialize Engines
    doc_converter = UniversalDocumentConverter()
    email_sender = None

    if not dry_run:
        try:
            provider = email_config.get("provider", "outlook") if email_config else "outlook"
            if provider == "smtp":
                if not email_config:
                    raise ValueError("SMTP configuration dictionary is missing")
                email_sender = SMTPSender(
                    host=email_config.get("smtp_host", ""),
                    port=int(email_config.get("smtp_port", 587)),
                    user=email_config.get("smtp_user", ""),
                    password=email_config.get("smtp_password", ""),
                    use_tls=email_config.get("use_tls"),
                    sender_name=email_config.get("sender_name")
                )
            else:
                email_sender = WinOutlookSender()
        except Exception as e:
            err_msg = f"❌ Email System Initialization Failed: {e}"
            log_callback(err_msg)
            return {"total": 0, "success": 0, "failed": 1, "records": [{"Status": err_msg}]}

    # 2. Load Dataset
    try:
        df = load_dataset(data_file)
    except Exception as e:
        err_msg = f"❌ Error loading dataset: {e}"
        log_callback(err_msg)
        return {"total": 0, "success": 0, "failed": 1, "records": [{"Status": err_msg}]}

    # 3. Handle Single-Row Test Mode
    if test_email:
        test_row = df.iloc[0:1].copy()
        test_row["Email"] = test_email
        df = test_row
        log_callback(f"🧪 Test Mode Active: Processing single test record for {test_email}")

    total_rows = len(df)
    log_callback(f"📊 Processing {total_rows} record(s)...")

    # 4. Phase 1: Compile Documents (.docx -> .pdf)
    compiled_items = []
    with UniversalDocumentConverter() as doc_converter:
        for idx, (original_idx, row) in enumerate(df.iterrows()):
            if cancel_event and cancel_event.is_set():
                log_callback("🛑 Job was cancelled by user during document compilation.")
                break

            email = str(row.get("Email", "")).strip()
            name = str(row.get("Name", row.get("Student_Name", row.get("Candidate_Name", f"Record_{idx+1}")))).strip()
            
            # Build replacement map
            replacements = {col: str(val).strip() for col, val in row.items()}
            
            # Resolve PDF filename
            resolved_filename = apply_template_replacements(pdf_filename_format, replacements)
            # Sanitize filename
            safe_filename = "".join(c for c in resolved_filename if c.isalnum() or c in (' ', '_', '-')).strip()
            if not safe_filename:
                safe_filename = f"Document_{idx+1}"

            word_path = os.path.join(pdf_folder, f"{safe_filename}.docx")
            pdf_path = os.path.join(pdf_folder, f"{safe_filename}.pdf")

            try:
                # Generate Word document
                doc = Document(template_file)
                replace_placeholders_in_doc(doc, replacements)
                doc.save(word_path)

                # Convert to PDF
                doc_converter.convert_to_pdf(word_path, pdf_path)
                
                # Remove intermediate Word file
                if os.path.exists(word_path):
                    try:
                        os.remove(word_path)
                    except Exception:
                        pass

                compiled_items.append({
                    "index": idx,
                    "name": name,
                    "email": email,
                    "pdf_path": pdf_path,
                    "replacements": replacements,
                    "row": row
                })
                log_callback(f"📄 [{idx+1}/{total_rows}] Compiled PDF for {name}")

            except Exception as e:
                err_msg = f"Conversion Error: {e}"
                log_callback(f"❌ [{idx+1}/{total_rows}] PDF Compilation Failed for {name}: {e}")
                log_records.append({
                    "Name": name,
                    "Email": email,
                    "Status": f"Failed: {err_msg}",
                    "Timestamp": datetime.now().isoformat()
                })

    # 5. Phase 2: Dispatch Emails (Parallel or Serial)
    def dispatch_single_email(item: Dict[str, Any]) -> Dict[str, Any]:
        if cancel_event and cancel_event.is_set():
            return {
                "Name": item["name"],
                "Email": item["email"],
                "Status": "Cancelled",
                "Timestamp": datetime.now().isoformat()
            }

        email = item["email"]
        name = item["name"]
        pdf_path = item["pdf_path"]
        replacements = item["replacements"]

        if not is_valid_email(email):
            log_callback(f"⚠️ [{item['index']+1}/{total_rows}] Skipped invalid email: {email}")
            return {
                "Name": name,
                "Email": email,
                "Status": "Failed: Invalid Email",
                "Timestamp": datetime.now().isoformat()
            }

        if dry_run or not email_sender:
            log_callback(f"📝 [{item['index']+1}/{total_rows}] Dry Run Completed for {name} ({email})")
            return {
                "Name": name,
                "Email": email,
                "Status": "Success (Dry Run)",
                "Timestamp": datetime.now().isoformat()
            }

        # Apply replacements to subject and body
        final_subject = apply_template_replacements(email_subject, replacements)
        final_body = apply_template_replacements(email_body, replacements)

        sent = False
        last_error = None
        for attempt in range(retries + 1):
            if cancel_event and cancel_event.is_set():
                break
            try:
                email_sender.send_email(email, final_subject, final_body, [pdf_path])
                sent = True
                break
            except Exception as e:
                last_error = e
                if attempt < retries:
                    time.sleep(delay)

        if sent:
            log_callback(f"✅ [{item['index']+1}/{total_rows}] Sent email to {email}")
            if delay > 0:
                time.sleep(delay)
            return {
                "Name": name,
                "Email": email,
                "Status": "Success",
                "Timestamp": datetime.now().isoformat()
            }
        else:
            log_callback(f"❌ [{item['index']+1}/{total_rows}] Delivery Failed to {email}: {last_error}")
            return {
                "Name": name,
                "Email": email,
                "Status": f"Failed: {last_error}",
                "Timestamp": datetime.now().isoformat()
            }

    # Dispatch Execution Strategy
    provider = email_config.get("provider", "outlook") if email_config else "outlook"
    use_threading = (provider == "smtp" and not dry_run and len(compiled_items) > 1)

    if use_threading:
        log_callback("🚀 Speed Mode: Parallel Asynchronous SMTP Dispatch (5 Workers)")
        with concurrent.futures.ThreadPoolExecutor(max_workers=5) as executor:
            future_to_item = {executor.submit(dispatch_single_email, it): it for it in compiled_items}
            for future in concurrent.futures.as_completed(future_to_item):
                try:
                    res = future.result()
                    log_records.append(res)
                except Exception as e:
                    item = future_to_item[future]
                    log_records.append({
                        "Name": item["name"],
                        "Email": item["email"],
                        "Status": f"Worker Error: {e}",
                        "Timestamp": datetime.now().isoformat()
                    })
    else:
        mode_label = "Dry Run" if dry_run else ("Outlook (Serial)" if provider == "outlook" else "SMTP (Serial)")
        log_callback(f"⚡ Execution Mode: {mode_label}")
        for it in compiled_items:
            if cancel_event and cancel_event.is_set():
                log_callback("🛑 Dispatch aborted by user.")
                break
            res = dispatch_single_email(it)
            log_records.append(res)

    # 6. Save Excel Log File
    try:
        new_log_df = pd.DataFrame(log_records)
        if os.path.exists(log_file_path):
            try:
                old_log_df = pd.read_excel(log_file_path)
                new_log_df = pd.concat([old_log_df, new_log_df], ignore_index=True)
            except Exception:
                pass
        new_log_df.to_excel(log_file_path, index=False)
        log_callback(f"📁 Execution log saved to {log_file_path}")
    except Exception as e:
        logger.warning(f"Could not write Excel log: {e}")

    # 7. Compute Summary
    success_count = sum(1 for r in log_records if "Success" in r.get("Status", ""))
    fail_count = sum(1 for r in log_records if "Failed" in r.get("Status", "") or "Error" in r.get("Status", ""))
    
    return {
        "total": len(log_records),
        "success": success_count,
        "failed": fail_count,
        "records": log_records
    }